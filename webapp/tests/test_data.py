# Tests für Presets, Geräteregistrierung und Zeitrekonstruktion
import csv
import glob
import io
import json
import time
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main
from src.core import data_manager as dm
from src.utils.constants import PRESETS_FILE, load_plant_presets

# Ohne "with": Der Lifespan und damit der BLE-Scanner starten nicht
client = TestClient(main.app)


def preset(**overrides):
    p = {"id": "test", "name": "Test", "target_soil_pct": 30,
         "sleep_multiplier": 1.2, "irrigation_start_delta": 15}
    p.update(overrides)
    return p


def test_invalid_presets_are_rejected():
    for bad in ({"target_soil_pct": 300}, {"target_soil_pct": 95}, {"sleep_multiplier": 0},
                {"irrigation_start_delta": 0}, {"irrigation_start_delta": 60}, {"name": ""}):
        assert client.put("/api/plant-presets", json=[preset(**bad)]).status_code == 422


def test_duplicate_presets_are_rejected():
    for presets in ([preset(), preset(name="Anders")], [preset(), preset(id="anders")]):
        assert client.put("/api/plant-presets", json=presets).status_code == 422


def test_device_gets_preset_values():
    assert client.put("/api/plant-presets", json=[preset()]).status_code == 200
    mac = "AA:BB:CC:DD:EE:01"
    client.post("/api/devices", json={"mac": mac, "name": "Kaktus",
                                      "preset_id": "test", "is_outdoor": False})
    config = client.get(f"/api/devices/{mac}").json()["config"]
    assert config["target_soil_pct"] == 30
    assert config["irrigation_start_delta"] == 15


def test_renaming_a_preset_keeps_the_device():
    assert client.put("/api/plant-presets", json=[preset()]).status_code == 200
    mac = "AA:BB:CC:DD:EE:03"
    client.post("/api/devices", json={"mac": mac, "name": "T", "preset_id": "test", "is_outdoor": False})
    assert client.put("/api/plant-presets", json=[preset(name='Kaktus 6" Topf')]).status_code == 200
    assert client.get(f"/api/devices/{mac}").json()["preset_id"] == "test"


def test_unknown_preset_is_rejected():
    body = {"mac": "AA:BB:CC:DD:EE:04", "name": "T", "preset_id": "gibt-es-nicht", "is_outdoor": False}
    assert client.post("/api/devices", json=body).status_code == 422


def test_device_type_is_validated():
    assert client.put("/api/plant-presets", json=[preset()]).status_code == 200
    mac = "AA:BB:CC:DD:EE:02"
    client.post("/api/devices", json={"mac": mac, "name": "T", "preset_id": "test", "is_outdoor": False})
    assert client.put(f"/api/devices/{mac}/settings", json={"type": "banane"}).status_code == 422
    assert client.put(f"/api/devices/{mac}/settings", json={"type": "outdoor"}).status_code == 200


def test_get_device_returns_a_copy():
    # Änderungen laufen nur über den DataManager (unter seinem Lock)
    mac = "AA:BB:CC:DD:EE:08"
    main.data_manager.add_or_update_device(mac, "Original", "test", False)
    main.data_manager.get_device(mac)["name"] = "Geändert"
    assert main.data_manager.get_device(mac)["name"] == "Original"
    assert main.data_manager.update_device(mac, notes="Notiz")
    assert main.data_manager.get_device(mac)["notes"] == "Notiz"
    assert not main.data_manager.update_device("AA:BB:CC:DD:EE:99", notes="x")


def test_invalid_mac_is_rejected():
    body = {"mac": "..\\..\\x", "name": "T", "preset_id": "test", "is_outdoor": False}
    assert client.post("/api/devices", json=body).status_code == 422


def test_preset_names_are_migrated_to_ids(tmp_path, monkeypatch):
    assert client.put("/api/plant-presets", json=[preset()]).status_code == 200
    settings_file = tmp_path / "device_settings.json"
    settings_file.write_text(json.dumps({
        "AA:BB:CC:DD:EE:05": {"name": "Alt", "plant_type": "Test"},
        "AA:BB:CC:DD:EE:06": {"name": "Weg", "plant_type": "Gelöschtes Preset"},
    }), encoding="utf-8")
    monkeypatch.setattr(dm, "SETTINGS_FILE", str(settings_file))

    manager = dm.DataManager()

    assert manager.get_device("AA:BB:CC:DD:EE:05")["preset_id"] == "test"
    assert manager.get_device("AA:BB:CC:DD:EE:06")["preset_id"] is None
    saved = json.loads(settings_file.read_text(encoding="utf-8"))
    assert all("plant_type" not in dev for dev in saved.values())


def test_unreadable_presets_are_kept():
    with open(PRESETS_FILE, "w", encoding="utf-8") as f:
        f.write("{kaputt")
    presets = load_plant_presets()
    assert presets[0]["id"] == "standard"
    backups = glob.glob(PRESETS_FILE + ".defekt-*")
    assert backups
    with open(backups[-1], encoding="utf-8") as f:
        assert f.read() == "{kaputt"


def awake(sleep_s, reset_reason=8):
    # 8 = ESP_RST_DEEPSLEEP
    return {"timestamp": 0, "reset_reason": reset_reason,
            "last_sleep_duration_s": sleep_s, "last_awake_time_ms": 2000}


def test_reconstruct_timestamps():
    before = int(time.time())
    buffered = [awake(3600), awake(1800)]
    online = awake(600)
    main.data_manager.reconstruct_timestamps(buffered, online)

    assert online["timestamp"] >= before
    # Jede Messung liegt Schlafdauer + Wachzeit des Nachfolgers davor
    assert buffered[1]["timestamp"] == int(online["timestamp"] - 600 - 2)
    assert buffered[0]["timestamp"] == int(buffered[1]["timestamp"] - 1800 - 2)


def test_reconstruct_stops_without_deep_sleep():
    # Nach einem normalen Reset ist der Abstand unbekannt
    buffered = [awake(3600)]
    online = awake(600, reset_reason=1)
    main.data_manager.reconstruct_timestamps(buffered, online)
    assert buffered[0]["timestamp"] is None


def test_excel_export_has_real_dates():
    assert client.put("/api/plant-presets", json=[preset()]).status_code == 200
    mac = "AA:BB:CC:DD:EE:07"
    client.post("/api/devices", json={"mac": mac, "name": "T", "preset_id": "test", "is_outdoor": False})
    main.data_manager.append_measurement(mac, {"timestamp": 1_790_000_000, "soil_pct": 40})
    main.data_manager.append_measurement(mac, {"timestamp": 0, "soil_pct": 41})

    resp = client.get(f"/api/devices/{mac}/export/excel")
    sheet = load_workbook(io.BytesIO(resp.content))["Messdaten"]

    assert sheet["B2"].value == datetime.fromtimestamp(1_790_000_000)
    assert sheet["B3"].value is None


def test_excel_export_leaves_missing_brain_decision_empty():
    # Über gepufferte Messungen hat das Brain nicht entschieden
    mac = "AA:BB:CC:DD:EE:0A"
    main.data_manager.add_or_update_device(mac, "T", "test", False)
    main.data_manager.append_measurement(mac, {"timestamp": 1_790_000_000, "soil_pct": 40, "online": False})
    main.data_manager.append_measurement(mac, {"timestamp": 1_790_003_600, "soil_pct": 39, "online": True,
                                               "brain_irrigate": False, "brain_sleep_s": 3600})

    resp = client.get(f"/api/devices/{mac}/export/excel")
    sheet = load_workbook(io.BytesIO(resp.content))["Messdaten"]
    column = [cell.value for cell in sheet[1]].index("Brain Bewässerung")

    assert [row[column].value for row in sheet.iter_rows(min_row=2)] == [None, "Nein"]


def test_old_csv_gets_new_columns():
    # Ältere Versionen schrieben die CSV ohne pump_dry
    mac = "AA:BB:CC:DD:EE:0B"
    path = dm._csv_path(mac)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[k for k in dm.CSV_FIELDS if k != "pump_dry"])
        writer.writeheader()
        writer.writerow({"timestamp": 1_790_000_000, "soil_pct": 40, "irrigated": True})

    main.data_manager.append_measurement(mac, {"timestamp": 1_790_003_600, "soil_pct": 41,
                                               "irrigated": True, "pump_dry": True})

    with open(path, newline="", encoding="utf-8") as f:
        assert next(csv.reader(f)) == dm.CSV_FIELDS
    history = main.data_manager.get_history(mac)
    assert [(r["soil_pct"], r["pump_dry"]) for r in history] == [(40, False), (41, True)]


def test_unknown_api_path_is_not_found():
    assert client.get("/api/gibt-es-nicht").status_code == 404
    # Andere Pfade sind Seiten des Dashboards
    assert client.get("/presets").status_code == 200
